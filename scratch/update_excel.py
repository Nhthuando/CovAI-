import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from copy import copy

def update_excel_backlog():
    excel_path = 'temp/C1SE.20_ProductBacklog-UserStory_Storelens_v1.1.xlsx'
    wb = openpyxl.load_workbook(excel_path)
    ws = wb['Product Backlog']

    print(f"Original max row: {ws.max_row}, max col: {ws.max_column}")

    # Row 5 Title check
    ws['B5'].value = "Product Backlog"

    # Define exact styles from row 8 template
    font_tnr_13 = Font(name='Times New Roman', size=13.0, bold=False)
    font_tnr_13_bold = Font(name='Times New Roman', size=13.0, bold=True)
    font_arial_10 = Font(name='Arial', size=10.0, bold=False)

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center_nowrap = Alignment(horizontal='center', vertical='center', wrap_text=False)

    thin_border = Border(
        left=Side(style='thin', color='FF000000'),
        right=Side(style='thin', color='FF000000'),
        top=Side(style='thin', color='FF000000'),
        bottom=Side(style='thin', color='FF000000')
    )

    items = [
        # Sprint 1 (336h)
        {
            "id": "PB-01",
            "heading": "Authentication & User Registration",
            "as_a": "Developer",
            "want": "Register a new account, log in securely with email and password, and manage session states with JWT tokens",
            "so_that": "I can securely access my private workspace, manage projects, and maintain protected project data",
            "criteria": "1. User can register with valid email, name, and strong password (>=8 chars, hashed with bcrypt).\n2. Login with valid credentials returns a signed JWT token and redirects to Dashboard.\n3. Invalid login displays clear error feedback without revealing account existence.\n4. Protected routes reject unauthenticated requests with HTTP 401 Unauthorized.\n5. Logout invalidates client session and redirects to login.",
            "remarks": "Express auth router, bcrypt, jsonwebtoken; AuthContext on client; User database model.",
            "priority": "5",
            "status_j": "Done",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 1",
            "estimate": 96,
            "status_o": "Done"
        },
        {
            "id": "PB-02",
            "heading": "User Profile & Encrypted GitHub Token Configuration",
            "as_a": "Developer",
            "want": "View my profile details and configure my GitHub Personal Access Token (PAT) securely",
            "so_that": "CovAI can authenticate against GitHub APIs to clone private repositories and setup webhooks without leaking credentials",
            "criteria": "1. User can input or update their GitHub Personal Access Token in Account Settings.\n2. Token is validated against GitHub API (/user) before being saved.\n3. Token is encrypted at rest in PostgreSQL using AES-256-GCM with server-side ENCRYPTION_KEY.\n4. Stored token is never exposed in plain text in client API responses or frontend state.",
            "remarks": "AES-256-GCM encryption with IV and authentication tag; GitHub REST API v3 integration.",
            "priority": "5",
            "status_j": "Done",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 1",
            "estimate": 72,
            "status_o": "Done"
        },
        {
            "id": "PB-03",
            "heading": "Project Management & Workspace Creation",
            "as_a": "Developer",
            "want": "Create, view, update, and delete software projects in my dedicated workspace",
            "so_that": "I can organize multiple repositories, inspect independent test histories, and maintain project-specific settings",
            "criteria": "1. User can create a new project with name, description, and source type (ZIP or GitHub).\n2. Dashboard lists all owned projects with status indicators, snapshot counts, and creation timestamps.\n3. Project details page displays project configuration, testing framework, and repository metadata.\n4. User can update project metadata or delete a project along with its cascaded snapshots and reports.\n5. Project access is strictly isolated by user ID via authMiddleware.",
            "remarks": "Express project controller, Prisma Project model, React workspace components.",
            "priority": "5",
            "status_j": "Done",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 1",
            "estimate": 96,
            "status_o": "Done"
        },
        {
            "id": "PB-04",
            "heading": "Snapshot Management & Source Versioning",
            "as_a": "Developer",
            "want": "Maintain immutable source-code snapshots for each uploaded archive or commit version",
            "so_that": "I can track coverage trends over time, compare code changes, and ensure reproducibility of quality evaluations",
            "criteria": "1. Each ingestion triggers creation of a new immutable snapshot record with unique SHA-256 hash.\n2. Snapshot stores version tag, commit message, source file tree, total files, and creation timestamp.\n3. Analysis results (coverage, CFG, AI test generation) are permanently linked to specific snapshot IDs.\n4. User can browse and toggle between past snapshots to view historical test reports.",
            "remarks": "Prisma Snapshot model, content-based SHA-256 hashing for deduplication and cache keys.",
            "priority": "5",
            "status_j": "Done",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 1",
            "estimate": 72,
            "status_o": "Done"
        },

        # Sprint 2 (336h)
        {
            "id": "PB-05",
            "heading": "Local ZIP Archive Project Ingestion",
            "as_a": "Developer",
            "want": "Upload project source code as a ZIP archive with automated secret sanitization and safety checks",
            "so_that": "I can analyze local codebases without committing to GitHub while protecting server security against malicious archives",
            "criteria": "1. System accepts .zip archives up to 200MB compressed size and max 1GB uncompressed expansion limit.\n2. Rejects zip bombs, recursive directory traversal (directory climbing ../), and invalid archive structures.\n3. Automatically sanitizes sensitive files (.env, .env.*, *.pem, *.key, id_rsa, node_modules).\n4. Extracts valid files to a temporary sandbox directory and uploads raw archive to Firebase Storage.\n5. Emits background progress events during upload, extraction, and verification.",
            "remarks": "Multer upload middleware, Adm-Zip / Archiver with safety guards, Firebase Storage bucket.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 2",
            "estimate": 64,
            "status_o": "Planned"
        },
        {
            "id": "PB-06",
            "heading": "GitHub Repository Ingestion & Branch Selection",
            "as_a": "Developer",
            "want": "Import a project repository directly from GitHub by selecting repository URL and branch",
            "so_that": "I can directly analyze hosted repositories without manual packaging or file transfers",
            "criteria": "1. User enters GitHub repository URL and branch name (defaults to main/master).\n2. System clones repository using stored encrypted user GitHub PAT or public clone URL.\n3. Performs depth=1 shallow clone to minimize latency and bandwidth consumption.\n4. Automatically strips sensitive configuration files (.env, .git/hooks) before creating snapshot.\n5. Stores commit hash, author, and commit message with the generated snapshot.",
            "remarks": "simple-git client, GitHub REST API, shallow clone sandbox worker.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 2",
            "estimate": 64,
            "status_o": "Planned"
        },
        {
            "id": "PB-07",
            "heading": "Automated Testing Framework Detection",
            "as_a": "Developer",
            "want": "Have the system automatically inspect project structure and detect configured testing frameworks",
            "so_that": "I don't have to manually configure test commands, config file paths, and coverage reporters",
            "criteria": "1. System parses package.json dependencies and devDependencies for Jest, Vitest, Mocha, Supertest, Playwright, Cypress.\n2. Detects configuration files: jest.config.*, vitest.config.*, playwright.config.*, cypress.config.*.\n3. Identifies test directory conventions (__tests__, *.test.js, *.spec.js, *.test.ts).\n4. Auto-selects appropriate test execution command with coverage flag (e.g., jest --coverage, vitest run --coverage).\n5. Displays detected framework, test runner version, and execution strategy on project overview.",
            "remarks": "FrameworkDetector service, AST-based config inspector, JSON schema validation.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 2",
            "estimate": 64,
            "status_o": "Planned"
        },
        {
            "id": "PB-08",
            "heading": "Isolated Docker Sandbox Test Execution",
            "as_a": "Developer",
            "want": "Execute automated test suites inside an isolated, resource-constrained Docker container",
            "so_that": "Tests execute securely without risking host compromise, dependency pollution, or resource exhaustion",
            "criteria": "1. Each test run spawns an ephemeral Docker container using a hardened Node.js base image.\n2. Enforces strict resource constraints: max 2 CPU cores, max 2GB RAM, execution timeout of 180 seconds.\n3. Disables network access or restricts to internal sandbox network during test run.\n4. Drops root privileges and runs execution under non-root unprivileged container user.\n5. Captures stdout, stderr, process exit code, and execution duration into persistent job logs.\n6. Container is automatically killed and pruned upon job completion or timeout expiry.",
            "remarks": "Docker Engine API / Dockerode, Docker Compose sandbox image, cgroups resource controls.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 2",
            "estimate": 80,
            "status_o": "Planned"
        },
        {
            "id": "PB-09",
            "heading": "Multi-Level Code Coverage Analysis",
            "as_a": "Developer",
            "want": "View comprehensive test coverage metrics across Statement, Branch, Function, and Line dimensions",
            "so_that": "I can precisely understand which source files and logic paths are thoroughly tested versus uncovered",
            "criteria": "1. Ingests and parses standard Istanbul/c8 JSON summary and coverage reports (coverage-final.json).\n2. Calculates and stores 4 standard coverage metrics: Statement %, Branch %, Function %, Line %.\n3. Presents overall project coverage score with color-coded badges (Green >= 80%, Yellow >= 60%, Red < 60%).\n4. Displays file-level coverage table with search, sorting, and line-by-line hit count highlighting.\n5. Flags uncovered lines, untested branch conditions (if-else, switch), and zero-call functions.",
            "remarks": "Istanbul/c8 report parser, CoverageService, Prisma CoverageReport & FileCoverage models.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 2",
            "estimate": 64,
            "status_o": "Planned"
        },

        # Sprint 3 (336h)
        {
            "id": "PB-10",
            "heading": "AST Parsing & Interactive Control Flow Graph (CFG)",
            "as_a": "Developer",
            "want": "Parse JavaScript/TypeScript ASTs and view interactive Control Flow Graphs for functions",
            "so_that": "I can visually understand execution branching, loop paths, and distinguish covered from uncovered paths",
            "criteria": "1. Uses Babel Parser (@babel/parser, @babel/traverse) to construct Abstract Syntax Trees for source files.\n2. Builds Control Flow Graphs mapping basic blocks, branching decisions (if, while, for, try/catch), and exit nodes.\n3. Correlates CFG nodes and edges with line/branch coverage data from the test execution.\n4. Renders interactive visual graph using Dagre / React Flow with green (covered) and red (uncovered) path markers.\n5. User can click a graph node to jump to the corresponding code line in the syntax-highlighted editor.",
            "remarks": "Babel parser, CFG generator service, Dagre layout engine, React Flow visualizer.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 3",
            "estimate": 64,
            "status_o": "Planned"
        },
        {
            "id": "PB-11",
            "heading": "Cyclomatic Complexity & Testing Gap Prioritization",
            "as_a": "Developer",
            "want": "Calculate Cyclomatic Complexity for every function and get a prioritized list of testing gaps",
            "so_that": "I can focus testing effort on high-risk, defect-prone functions that have high complexity and low coverage",
            "criteria": "1. Calculates McCabe Cyclomatic Complexity (CC = E - N + 2P) for each function and module.\n2. Computes Risk Score based on formula combining High Complexity (CC > 10) and Low Branch Coverage (< 60%).\n3. Generates prioritized 'Testing Gaps' list sorted by risk severity (Critical, High, Medium, Low).\n4. Provides actionable gap explanation: identifies specific uncovered branches and potential edge cases.\n5. Persists complexity metrics and gap records associated with the snapshot.",
            "remarks": "Static analysis service, McCabe formula implementation, Risk ranking algorithm.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 3",
            "estimate": 56,
            "status_o": "Planned"
        },
        {
            "id": "PB-12",
            "heading": "Multi-Framework Test Runner Support",
            "as_a": "QA Engineer",
            "want": "Execute tests across Unit, Integration, and E2E frameworks (Jest, Vitest, Supertest, Playwright, Cypress)",
            "so_that": "I can evaluate test coverage and execution health across all layers of the testing pyramid",
            "criteria": "1. System supports execution adapters for Jest, Vitest, Supertest, Playwright, and Cypress.\n2. Executes headless unit/integration test suites and captures standard JUnit/TAP/JSON test summaries.\n3. Normalizes test count (passed, failed, skipped), duration, and coverage output across different frameworks.\n4. Displays consolidated test result breakdown by test suite and test file in the QA dashboard.\n5. Gracefully handles framework-specific error outputs and missing reporter configurations.",
            "remarks": "Multi-framework runner adapters, Docker headless browsers for Playwright/Cypress.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 3",
            "estimate": 48,
            "status_o": "Planned"
        },
        {
            "id": "PB-13",
            "heading": "AI-Assisted Test Skeleton Generation via Gemini API",
            "as_a": "Developer",
            "want": "Request AI-generated test skeletons for uncovered functions using Google Gemini API",
            "so_that": "I can quickly get scaffolded test structures, describe blocks, and mock setups for untested functions",
            "criteria": "1. Developer selects an uncovered or low-coverage function from the Testing Gaps list.\n2. System extracts bounded context: function signature, docstrings, imports, dependencies, and uncovered branches.\n3. Sends engineered prompt to Google Gemini API (gemini-1.5-pro / gemini-1.5-flash) enforcing strict test formatting.\n4. Generates clean test skeleton with describe, it/test blocks, mock stubs, and descriptive test scenario names.\n5. Fallback mechanism provides rule-based test template if AI service is temporarily unreachable or quota exceeded.",
            "remarks": "Google Gemini API client, Prompt engineering templates, Token limiter & quota tracker.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 3",
            "estimate": 56,
            "status_o": "Planned"
        },
        {
            "id": "PB-14",
            "heading": "AI-Assisted Full Runnable Test Generation",
            "as_a": "Developer",
            "want": "Generate complete, runnable unit test implementations targeting specific uncovered branches",
            "so_that": "I can directly run generated tests to close coverage gaps without writing boilerplate setup manually",
            "criteria": "1. Developer requests full test generation for target function with specified branch focus.\n2. System provides Gemini with source code, existing test mock patterns, and branch condition details.\n3. Generated code includes: imports, input data fixtures, mocked dependencies, assertions, and boundary checks.\n4. Output is parsed, sanitized, and verified for syntactic correctness using Babel parse check.\n5. Stores generation history, prompt token count, model version, and generation metadata.",
            "remarks": "Gemini API integration, AST syntax validator, Test generator service.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 3",
            "estimate": 64,
            "status_o": "Planned"
        },
        {
            "id": "PB-15",
            "heading": "Generated Test Case Validation & Code Export",
            "as_a": "Developer",
            "want": "Review, edit, validate in sandbox, and export AI-generated test files into my project",
            "so_that": "I can verify test correctness and safely integrate new test suites into my codebase",
            "criteria": "1. Split-screen UI displays original source code alongside AI-generated test code with syntax highlighting.\n2. Developer can edit generated test code directly in the web editor before accepting.\n3. 'Run in Sandbox' button executes the new test in an isolated Docker container against the project source.\n4. Reports test execution status (Passed/Failed) and new coverage delta achieved.\n5. User can download test file (*.test.js) or copy code to clipboard with one click.",
            "remarks": "Monaco Editor / PrismJS integration, Sandbox test validator, Coverage diff calculator.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 3",
            "estimate": 48,
            "status_o": "Planned"
        },

        # Sprint 4 (144h)
        {
            "id": "PB-16",
            "heading": "GitHub Webhook Automation & Continuous Analysis",
            "as_a": "Developer",
            "want": "Configure GitHub repository webhooks to automatically trigger test runs upon new commits or PRs",
            "so_that": "My project coverage and quality scores are continuously analyzed without manual trigger actions",
            "criteria": "1. System generates a unique webhook endpoint URL and HMAC-SHA256 secret for each GitHub project.\n2. Listens for GitHub push and pull_request events and verifies payload signature (X-Hub-Signature-256).\n3. Automatically creates a new project snapshot with commit metadata and queues analysis job in BullMQ.\n4. Posts commit status check back to GitHub API (Pending, Success, Failure) with link to CovAI report.\n5. User can enable, disable, or regenerate webhook secret in Project Settings.",
            "remarks": "Express webhook handler, crypto HMAC verification, GitHub Status API.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 4",
            "estimate": 40,
            "status_o": "Planned"
        },
        {
            "id": "PB-17",
            "heading": "Asynchronous BullMQ Queue & Worker Monitoring",
            "as_a": "Administrator",
            "want": "Manage and monitor BullMQ background job queues and worker execution states in real-time",
            "so_that": "I can ensure smooth asynchronous analysis, detect queue bottlenecks, and retry failed analysis tasks",
            "criteria": "1. Background tasks (ZIP extraction, Docker test runs, AST parsing, AI generation) are queued in BullMQ with Redis.\n2. Workers process jobs with configured concurrency, timeouts, and exponential backoff retry policies.\n3. Admin can view queue metrics: Active, Waiting, Completed, Failed, and Delayed job counts.\n4. Admin can inspect job error stack traces, view execution logs, and trigger manual retries for failed jobs.\n5. System gracefully handles Redis disconnections and logs worker health metrics.",
            "remarks": "BullMQ, Redis client (ioredis), Bull-Board dashboard or custom admin queue inspector.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 4",
            "estimate": 40,
            "status_o": "Planned"
        },
        {
            "id": "PB-18",
            "heading": "Real-Time WebSocket Notifications",
            "as_a": "Developer",
            "want": "Receive real-time WebSocket push notifications when long-running analysis or AI generation completes",
            "so_that": "I don't have to repeatedly refresh the browser or wait idle on the screen during background processing",
            "criteria": "1. Client establishes authenticated Socket.IO connection upon logging in.\n2. Server emits real-time events on job progress: job:progress, job:completed, job:failed.\n3. UI displays non-intrusive toast notifications when coverage analysis, AST/CFG, or AI test generation finishes.\n4. Notifications dropdown menu stores recent notification history with direct links to generated reports.\n5. Automatically re-syncs state if socket reconnects after network drop.",
            "remarks": "Socket.IO server & client, Redis pub/sub adapter, React notification toast component.",
            "priority": "3",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 4",
            "estimate": 32,
            "status_o": "Planned"
        },
        {
            "id": "PB-19",
            "heading": "Software Quality Overview Dashboard & Radar Metrics",
            "as_a": "QA Engineer",
            "want": "View a centralized Quality Dashboard with composite quality scores and multi-dimensional radar charts",
            "so_that": "I can holistically evaluate software release readiness across coverage, complexity, maintainability, and security",
            "criteria": "1. Computes Composite Quality Score (0-100) combining Coverage (40%), Complexity (20%), Security (20%), Maintainability (20%).\n2. Renders interactive Radar Chart visual comparing the project's quality dimensions against target benchmarks.\n3. Displays summary metric cards: Total Tests, Pass Rate %, Overall Coverage %, Total Functions, Critical Gaps.\n4. Provides quick action links to drill down into file details, test reports, and AI generation tools.\n5. Supports exporting summary overview as PNG or formatted PDF.",
            "remarks": "Recharts / Chart.js radar charts, QualityScoringService, React Dashboard UI.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 4",
            "estimate": 32,
            "status_o": "Planned"
        },

        # Sprint 5 (336h)
        {
            "id": "PB-20",
            "heading": "AI Evaluation Knowledge Ingestion & Document Parsing",
            "as_a": "AI Evaluator",
            "want": "Upload project business documents, API specifications, FAQ files, and operational rule sets",
            "so_that": "CovAI can establish the project ground truth required for evaluating internal chatbot trustworthiness",
            "criteria": "1. Accepts PDF, Markdown, text, and OpenAPI/Swagger JSON/YAML files representing project knowledge.\n2. Extracts and parses text, section hierarchies, API endpoints, parameters, and business constraints.\n3. Sanitizes uploaded knowledge files and checks file size limits (max 50MB per document package).\n4. Stores ingested documents in structured knowledge base linked to the project snapshot.\n5. Displays document preview, parsed token counts, and indexing status in the Evaluation workspace.",
            "remarks": "Document parser (pdf-parse, markdown-it, swagger-parser), Prisma KnowledgeDoc model.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 5",
            "estimate": 48,
            "status_o": "Planned"
        },
        {
            "id": "PB-21",
            "heading": "Evaluation Profile Builder & Benchmark Scenario Generation",
            "as_a": "AI Evaluator",
            "want": "Configure target chatbot endpoints and automatically generate candidate evaluation scenarios",
            "so_that": "I can create a structured benchmark dataset covering Knowledge Q&A, Usage Guidance, and Action/Tool execution",
            "criteria": "1. User configures chatbot adapter: HTTP POST endpoint, authentication headers, request/response JSON schema mapping.\n2. AI Profile Builder analyzes ingested documents and proposes candidate test scenarios categorized by task type:\n   - Knowledge Q&A scenarios (factual retrieval, policy clarification, out-of-scope refusal).\n   - Usage Guidance scenarios (step-by-step UI navigation, button clicks, workflows).\n   - Action/Tool scenarios (API calls, data mutations, query execution).\n3. Generates candidate prompts, expected answers, expected tool calls, and ground truth citations.\n4. Persists evaluation profile with versioning and configuration parameters.",
            "remarks": "EvaluationProfileBuilder service, Gemini scenario generator, JSON schema validator.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 5",
            "estimate": 56,
            "status_o": "Planned"
        },
        {
            "id": "PB-22",
            "heading": "Domain Expert Rule & Ground Truth Confirmation",
            "as_a": "Domain Expert / Evaluator",
            "want": "Review, edit, approve, or add ground truth facts, expected UI sequences, and forbidden actions",
            "so_that": "The evaluation benchmark precisely reflects validated business logic and prevents false positive penalties",
            "criteria": "1. Interactive scenario management table displaying prompt, expected behavior, ground truth facts, and rules.\n2. Expert can edit prompt text, adjust expected answer keywords, and configure required citations.\n3. For Action/Tool tasks: expert can specify mandatory confirmation prompts, allowed tools, and strictly forbidden actions.\n4. Expert marks scenarios as 'Approved for Benchmark' or 'Excluded'.\n5. Tracks author, edit history, and approval timestamp for complete auditability.",
            "remarks": "Ground truth verification UI, Role-based editing permissions, Prisma EvaluationScenario model.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 5",
            "estimate": 40,
            "status_o": "Planned"
        },
        {
            "id": "PB-23",
            "heading": "Multi-Task Chatbot Evaluation Runner",
            "as_a": "AI Evaluator",
            "want": "Execute automated benchmark evaluation runs against the target chatbot interface",
            "so_that": "The system systematically tests all approved scenarios and captures raw responses, latency, and metadata",
            "criteria": "1. Evaluation runner iterates through all approved scenarios in the active evaluation profile.\n2. Dispatches prompts to target chatbot endpoint with session context, timeout limits, and rate-limiting throttling.\n3. Collects raw text response, citations, generated tool calls, HTTP status, and response latency (ms).\n4. Executes multi-turn conversations for context-retention and clarification scenarios.\n5. Records complete execution trace and persists evaluation run metadata in database.",
            "remarks": "EvaluationRunner service, Axios HTTP client with retry & timeout, BullMQ background job.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 5",
            "estimate": 56,
            "status_o": "Planned"
        },
        {
            "id": "PB-24",
            "heading": "Chatbot Tool & Action Execution Verification",
            "as_a": "AI Evaluator",
            "want": "Verify tool selection accuracy, parameter schemas, safety confirmations, and forbidden actions",
            "so_that": "I can prevent risky external mutations and ensure the chatbot safely adheres to API authorization rules",
            "criteria": "1. Verifies that chatbot selected the exact expected tool/API endpoint for action-oriented prompts.\n2. Validates extracted parameters against expected schema types, boundary values, and mandatory fields.\n3. Verifies that high-risk actions (delete, transfer, checkout) require explicit user confirmation before execution.\n4. Immediately flags safety violations if chatbot attempts forbidden or unauthorized tool actions.\n5. Checks for duplicate tool executions or unhandled error loops.",
            "remarks": "ActionVerifier engine, JSON schema validation, Tool trace audit inspector.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 5",
            "estimate": 48,
            "status_o": "Planned"
        },
        {
            "id": "PB-25",
            "heading": "3-Layer Trustworthiness Multi-Dimensional Scoring",
            "as_a": "AI Evaluator",
            "want": "Evaluate chatbot outputs across the 3-Layer framework (Core Criteria, Task Module, Project Rules)",
            "so_that": "I receive objective, multi-dimensional scores on Correctness, Faithfulness, Hallucination, Safety, and Latency",
            "criteria": "1. Evaluates Layer 1 Core Criteria: Correctness %, Faithfulness %, Relevance %, Hallucination Rate %, Appropriate Refusal %, Safety %, Latency.\n2. Evaluates Layer 2 Task Module Criteria:\n   - Q&A: Source accuracy, citation presence.\n   - Guidance: UI step correctness, navigation completeness.\n   - Action: Tool selection accuracy %, parameter precision %, confirmation compliance %.\n3. Evaluates Layer 3 Project Rules: Compliance with specific verified ground-truth constraints and policies.\n4. Computes Overall Trustworthiness Index (0-100) using weighted multi-criteria aggregation.\n5. Stores granular sub-scores and violation tags for every evaluated scenario.",
            "remarks": "3-Layer Scoring Engine, LLM-as-a-judge evaluation prompts with deterministic heuristics.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 5",
            "estimate": 48,
            "status_o": "Planned"
        },
        {
            "id": "PB-26",
            "heading": "Chatbot Diagnosis, Benchmark Comparison & Report Export",
            "as_a": "AI Evaluator",
            "want": "Inspect detailed failure diagnosis, compare multiple chatbot versions/models, and export benchmark reports",
            "so_that": "I can identify root causes of chatbot errors, track model improvements, and present compliance evidence",
            "criteria": "1. Displays diagnostic report: lists failed scenarios, hallucination snippets, parameter mismatches, and risk levels.\n2. Generates automated AI improvement suggestions for prompts, retrieval contexts, and guardrails.\n3. Supports side-by-side benchmark comparison between multiple chatbot versions or configurations under same profile.\n4. Visualizes comparison charts: Trustworthiness score delta, latency trends, and error category distribution.\n5. Exports comprehensive evaluation reports in structured PDF and JSON formats.",
            "remarks": "ComparisonEngine, Diagnostic reporter, PDF/JSON report exporter.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 5",
            "estimate": 40,
            "status_o": "Planned"
        },

        # Sprint 6 (168h)
        {
            "id": "PB-27",
            "heading": "End-to-End System Integration & Regression Testing",
            "as_a": "QA Engineer",
            "want": "Perform automated end-to-end integration and regression test runs across the complete platform pipeline",
            "so_that": "I can ensure all interconnected modules (Ingestion -> Sandbox -> Coverage -> CFG -> AI -> Evaluation) work seamlessly",
            "criteria": "1. Executes automated integration test suite covering end-to-end workflow: user login, project import, test execution, coverage report generation, AST/CFG display, AI test generation, and chatbot evaluation.\n2. Verifies data integrity across PostgreSQL database tables, Redis queues, and Firebase storage artifacts.\n3. Validates error recovery when Docker daemon fails, sandbox times out, or target chatbot API returns HTTP 500.\n4. Generates automated test execution logs and defect reports for any discovered regressions.",
            "remarks": "Supertest API testing suite, Jest integration tests, Mock Docker/Gemini testing fixtures.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 6",
            "estimate": 96,
            "status_o": "Planned"
        },
        {
            "id": "PB-28",
            "heading": "Security Hardening, Quotas & Container Sandboxing Verification",
            "as_a": "Administrator",
            "want": "Conduct security penetration checks, rate-limit verifications, and Docker sandbox escape audits",
            "so_that": "The platform is robustly defended against malicious code execution, token depletion, and unauthorized data access",
            "criteria": "1. Validates that API endpoints enforce rate limiting (express-rate-limit) and reject brute force attacks.\n2. Verifies that encrypted GitHub PATs and user credentials cannot be decrypted without server secret key.\n3. Audits Docker sandbox isolation: tests container escape attempts, fork bomb resistance, and memory exhaustion limits.\n4. Verifies that Gemini AI API token usage caps prevent unexpected billing or quota exhaustion.\n5. Produces security compliance audit log ready for Capstone 1 presentation.",
            "remarks": "Security audit suite, OWASP Top 10 verification, Docker security profiling.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Non-Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 6",
            "estimate": 72,
            "status_o": "Planned"
        },

        # Sprint 7 (264h)
        {
            "id": "PB-29",
            "heading": "Production Deployment & Docker Compose Multi-Container Orchestration",
            "as_a": "Administrator",
            "want": "Deploy the complete CovAI platform using Docker Compose multi-container orchestration",
            "so_that": "The client, server, Redis, and PostgreSQL services run reliably in production for defense demo and user trials",
            "criteria": "1. Docker Compose configuration orchestrates client (Vite/Nginx), server (Express API), database (PostgreSQL), and Redis.\n2. Environment variables are securely injected from .env with validation on startup.\n3. Database migrations run automatically via prisma migrate deploy upon server startup.\n4. Nginx reverse proxy routes client traffic, API endpoints (/api), and WebSocket connections (/socket.io).\n5. System starts cleanly with docker compose up --build with healthcheck passing on all services.",
            "remarks": "Docker Compose, Nginx reverse proxy, multi-stage Dockerfiles, Prisma migration pipeline.",
            "priority": "5",
            "status_j": "Planned",
            "type": "Non-Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 7",
            "estimate": 144,
            "status_o": "Planned"
        },
        {
            "id": "PB-30",
            "heading": "Admin Platform Oversight, Audit Logs & API Quota Management",
            "as_a": "Administrator",
            "want": "Monitor system-wide platform health, view audit logs, and configure global Gemini AI token limits",
            "so_that": "I can ensure operational stability, track user usage patterns, and prevent unexpected infrastructure expenses",
            "criteria": "1. Admin dashboard displays active users, total projects analyzed, Docker container runtime stats, and error rates.\n2. Real-time tracking of Google Gemini API token usage per project and aggregate monthly consumption.\n3. Configurable rate limits and daily token quotas per user tier to prevent abuse.\n4. Audit log records all critical events: user logins, project deletions, webhook triggers, and permission changes.\n5. Admin can export platform usage metrics and audit logs for academic review and grading.",
            "remarks": "Admin dashboard, AuditLog service, Gemini token counter & quota enforcement middleware.",
            "priority": "4",
            "status_j": "Planned",
            "type": "Functional",
            "proj_id": "COVAI",
            "sprint": "Sprint 7",
            "estimate": 120,
            "status_o": "Planned"
        }
    ]

    # Clear old data rows from 8 onwards
    for r in range(8, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.border = Border()
            cell.fill = PatternFill(fill_type=None)

    # Populate 30 items
    start_row = 8
    for idx, item in enumerate(items):
        r = start_row + idx
        ws.row_dimensions[r].height = 95.0

        # Col B (2): Id
        cell = ws.cell(row=r, column=2, value=item["id"])
        cell.font = font_tnr_13
        cell.alignment = align_center
        cell.border = thin_border

        # Col C (3): Heading
        cell = ws.cell(row=r, column=3, value=item["heading"])
        cell.font = font_tnr_13
        cell.alignment = align_left_wrap
        cell.border = thin_border

        # Col D (4): As a..
        cell = ws.cell(row=r, column=4, value=item["as_a"])
        cell.font = font_tnr_13
        cell.alignment = align_left_wrap
        cell.border = thin_border

        # Col E (5): I want to ..
        cell = ws.cell(row=r, column=5, value=item["want"])
        cell.font = font_tnr_13
        cell.alignment = align_left_wrap
        cell.border = thin_border

        # Col F (6): so that ..
        cell = ws.cell(row=r, column=6, value=item["so_that"])
        cell.font = font_tnr_13
        cell.alignment = align_left_wrap
        cell.border = thin_border

        # Col G (7): Acceptance Criteria
        cell = ws.cell(row=r, column=7, value=item["criteria"])
        cell.font = font_tnr_13
        cell.alignment = align_left_wrap
        cell.border = thin_border

        # Col H (8): Remarks
        cell = ws.cell(row=r, column=8, value=item["remarks"])
        cell.font = font_arial_10
        cell.alignment = align_left_wrap
        cell.border = thin_border

        # Col I (9): Priority
        cell = ws.cell(row=r, column=9, value=item["priority"])
        cell.font = font_tnr_13
        cell.alignment = align_center
        cell.border = thin_border

        # Col J (10): Status
        cell = ws.cell(row=r, column=10, value=item["status_j"])
        cell.font = font_arial_10
        cell.alignment = align_center
        cell.border = thin_border

        # Col K (11): Type
        cell = ws.cell(row=r, column=11, value=item["type"])
        cell.font = font_arial_10
        cell.alignment = align_center
        cell.border = thin_border

        # Col L (12): Project Id
        cell = ws.cell(row=r, column=12, value=item["proj_id"])
        cell.font = font_arial_10
        cell.alignment = align_center
        cell.border = thin_border

        # Col M (13): Sprint No
        cell = ws.cell(row=r, column=13, value=item["sprint"])
        cell.font = font_tnr_13
        cell.alignment = align_center
        cell.border = thin_border

        # Col N (14): Estimate (Hours)
        cell = ws.cell(row=r, column=14, value=item["estimate"])
        cell.font = font_arial_10
        cell.alignment = align_center
        cell.border = thin_border

        # Col O (15): Status
        cell = ws.cell(row=r, column=15, value=item["status_o"])
        cell.font = font_tnr_13
        cell.alignment = align_center
        cell.border = thin_border

    wb.save(excel_path)
    print(f"Successfully populated {len(items)} items into {excel_path}")

if __name__ == '__main__':
    update_excel_backlog()
