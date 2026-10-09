import json
import hashlib
import zipfile
import io
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from openpyxl import load_workbook

ROOT = Path(r"D:\HuuThuan - Project\NCKH\CovAI")
TEMP = ROOT / "temp"
SCRATCH = ROOT / "scratch"
SCRATCH.mkdir(exist_ok=True)

FILES = {
    "proposal_pdf": TEMP / "C1SE.30_Proposal_ver1.2.pdf",
    "plan_pdf": TEMP / "C1SE.30_ProjectPlan_ver_1.1-C1SE.30_Sprint1_COVAI_ver_1.0.pdf",
    "proposal_docx": TEMP / "C1SE.30_Proposal_ver1.2.docx",
    "backlog_docx": TEMP / "C1SE.30_ProductBacklog_CovAI_ver1.1.docx",
    "backlog_xlsx": TEMP / "C1SE.20_ProductBacklog-UserStory_Storelens_v1.1.xlsx",
    "backlog_docx_bak": TEMP / "C1SE.30_ProductBacklog_CovAI_ver1.1.docx.bak",
    "backlog_xlsx_bak": TEMP / "C1SE.20_ProductBacklog-UserStory_Storelens_v1.1.xlsx.bak",
}


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def pdf_info(path, text_path):
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        pages.append({"page": i, "chars": len(text), "text": text})
    text_path.write_text("\n\n===== PAGE %d =====\n" % 1 + pages[0]["text"] if pages else "", encoding="utf-8")
    if len(pages) > 1:
        with text_path.open("a", encoding="utf-8") as f:
            for p in pages[1:]:
                f.write("\n\n===== PAGE %d =====\n%s" % (p["page"], p["text"]))
    return {
        "path": str(path),
        "sha256": sha(path),
        "pages": len(pages),
        "page_chars": [{"page": p["page"], "chars": p["chars"]} for p in pages],
        "metadata": {str(k): str(v) for k, v in (reader.metadata or {}).items()},
    }


def fmt_run(run):
    font = run.font
    return {
        "text": run.text,
        "bold": font.bold,
        "italic": font.italic,
        "underline": str(font.underline) if font.underline is not None else None,
        "size_pt": float(font.size.pt) if font.size else None,
        "name": font.name,
        "color": str(font.color.rgb) if font.color and font.color.type == "rgb" else None,
    }


def para_record(p):
    pf = p.paragraph_format
    return {
        "text": p.text,
        "style": p.style.name if p.style else None,
        "alignment": str(p.alignment),
        "left_indent": float(pf.left_indent.inches) if pf.left_indent else None,
        "right_indent": float(pf.right_indent.inches) if pf.right_indent else None,
        "first_line_indent": float(pf.first_line_indent.inches) if pf.first_line_indent else None,
        "space_before_pt": float(pf.space_before.pt) if pf.space_before else None,
        "space_after_pt": float(pf.space_after.pt) if pf.space_after else None,
        "line_spacing": pf.line_spacing,
        "runs": [fmt_run(r) for r in p.runs],
    }


def table_record(table):
    rows = []
    for row in table.rows:
        rows.append([[para_record(p) for p in cell.paragraphs] for cell in row.cells])
    return {"rows": len(table.rows), "cols": len(table.columns), "data": rows}


def part_records(container):
    out = {"paragraphs": [para_record(p) for p in container.paragraphs], "tables": [table_record(t) for t in container.tables]}
    return out


def docx_info(path, full=False):
    doc = Document(str(path))
    sections = []
    for s in doc.sections:
        sections.append({
            "page_width_in": float(s.page_width.inches),
            "page_height_in": float(s.page_height.inches),
            "top_margin_in": float(s.top_margin.inches),
            "bottom_margin_in": float(s.bottom_margin.inches),
            "left_margin_in": float(s.left_margin.inches),
            "right_margin_in": float(s.right_margin.inches),
            "header_distance_in": float(s.header_distance.inches),
            "footer_distance_in": float(s.footer_distance.inches),
            "different_first_page_header_footer": s.different_first_page_header_footer,
            "header": part_records(s.header),
            "footer": part_records(s.footer),
        })
    styles = []
    for st in doc.styles:
        if st.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        pf = st.paragraph_format
        font = st.font
        styles.append({
            "name": st.name,
            "base_style": st.base_style.name if st.base_style else None,
            "font_name": font.name,
            "font_size_pt": float(font.size.pt) if font.size else None,
            "bold": font.bold,
            "italic": font.italic,
            "color": str(font.color.rgb) if font.color and font.color.type == "rgb" else None,
            "alignment": str(pf.alignment),
            "space_before_pt": float(pf.space_before.pt) if pf.space_before else None,
            "space_after_pt": float(pf.space_after.pt) if pf.space_after else None,
            "line_spacing": pf.line_spacing,
            "left_indent": float(pf.left_indent.inches) if pf.left_indent else None,
            "first_line_indent": float(pf.first_line_indent.inches) if pf.first_line_indent else None,
        })
    package = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            package.append({"name": info.filename, "size": info.file_size, "crc": info.CRC})
    result = {
        "path": str(path),
        "sha256": sha(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "section_count": len(doc.sections),
        "sections": sections,
        "styles": styles,
        "package": package,
        "body": part_records(doc),
    }
    if not full:
        result["body"] = {
            "paragraphs": [p for p in result["body"]["paragraphs"] if p["text"].strip()],
            "tables": [{"rows": t["rows"], "cols": t["cols"], "data": [[[q["text"] for q in cell] for cell in row] for row in t["data"]]} for t in result["body"]["tables"]],
        }
    return result


def xlsx_info(path):
    # openpyxl rejects the .bak extension even though the backup is a valid
    # OOXML workbook, so load it from an in-memory stream for read-only QA.
    wb = load_workbook(io.BytesIO(path.read_bytes()), data_only=False, read_only=False)
    sheets = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            vals = [c.value for c in row]
            if any(v is not None for v in vals):
                cells.append({"row": row[0].row, "values": vals})
        sheets.append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "sheet_view_show_gridLines": ws.sheet_view.showGridLines,
            "print_area": str(ws.print_area) if ws.print_area else None,
            "print_title_rows": str(ws.print_title_rows) if ws.print_title_rows else None,
            "print_title_cols": str(ws.print_title_cols) if ws.print_title_cols else None,
            "tables": {name: str(tbl.ref) for name, tbl in ws.tables.items()},
            "nonempty_rows": cells,
        })
    return {"path": str(path), "sha256": sha(path), "sheet_count": len(wb.worksheets), "sheets": sheets}


def main():
    result = {"files": {k: {"path": str(v), "exists": v.exists(), "sha256": sha(v) if v.exists() else None, "size": v.stat().st_size if v.exists() else None} for k, v in FILES.items()}}
    result["proposal_pdf"] = pdf_info(FILES["proposal_pdf"], SCRATCH / "proposal_pdf_text.txt")
    result["plan_pdf"] = pdf_info(FILES["plan_pdf"], SCRATCH / "plan_pdf_text.txt")
    result["proposal_docx"] = docx_info(FILES["proposal_docx"], full=False)
    result["backlog_docx"] = docx_info(FILES["backlog_docx"], full=True)
    result["backlog_xlsx"] = xlsx_info(FILES["backlog_xlsx"])
    result["backlog_docx_bak"] = docx_info(FILES["backlog_docx_bak"], full=False)
    result["backlog_xlsx_bak"] = xlsx_info(FILES["backlog_xlsx_bak"])
    (SCRATCH / "source_analysis.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({k: {"pages": v.get("pages"), "paragraph_count": v.get("paragraph_count"), "table_count": v.get("table_count"), "sheet_count": v.get("sheet_count")} for k, v in result.items() if isinstance(v, dict) and k not in {"files"}}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
