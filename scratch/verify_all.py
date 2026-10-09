import docx
import openpyxl

def verify_all():
    print("=" * 60)
    print("STARTING COMPREHENSIVE VERIFICATION OF COVAI PRODUCT BACKLOG DELIVERABLES")
    print("=" * 60)

    # 1. Verify Word Document
    docx_path = 'temp/C1SE.30_ProductBacklog_CovAI_ver1.1.docx'
    doc = docx.Document(docx_path)

    full_docx_text = ' '.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    hien_found = 'Hien' in full_docx_text or 'Nguyen Thanh Long' in full_docx_text
    print(f"\n[DOCX] Leaked member 'Hien' found: {hien_found} (Expected: False)")
    assert not hien_found, "Error: Leaked member Hien still found in docx!"

    # Verify Cover page members
    cover_names = [doc.paragraphs[i].text.strip() for i in [12, 13, 14, 15]]
    print(f"[DOCX] Cover page submitted by: {cover_names}")
    expected_members = ['Dinh, Huynh Tan', 'Phuc, Tran Huu', 'Quy, Nguyen Duy', 'Thuan, Ngo Huu']
    assert cover_names == expected_members, f"Error: Cover page members do not match {expected_members}"

    # Verify Table 0 Project Information
    t0 = doc.tables[0]
    print(f"[DOCX] Table 0 (Project Info) rows: {len(t0.rows)}")
    t0_members = [t0.rows[r].cells[1].text.strip() for r in [9, 10, 11, 12]]
    print(f"[DOCX] Table 0 Team members: {t0_members}")
    assert t0_members == expected_members, "Error: Table 0 team members mismatch!"

    # Verify Table 1 Approvals
    t1 = doc.tables[1]
    print(f"[DOCX] Table 1 (Approvals) rows: {len(t1.rows)}")
    assert len(t1.rows) == 4, f"Error: Table 1 should have 4 rows, found {len(t1.rows)}"

    # Verify Tables count and Backlog Specification
    print(f"[DOCX] Total tables in docx: {len(doc.tables)} (Expected: 11)")
    # Tables are:
    # 0: Project Info
    # 1: Approvals
    # 2: Document Details
    # 3: Revision History
    # 4: References
    # 5: Admin Backlog (PB01-PB05)
    # 6: Developer Backlog (PB01-PB11)
    # 7: QA Engineer Backlog (PB01-PB07)
    # 8: AI Evaluator Backlog (PB01-PB08) (Newly added!)
    # 9: Constraints
    # 10: Stakeholders Summary
    assert len(doc.tables) == 11, f"Expected 11 tables, got {len(doc.tables)}"

    t5_admin = doc.tables[5]
    print(f"[DOCX] Admin Backlog items: {len(t5_admin.rows) - 1} (Expected: 5)")
    assert len(t5_admin.rows) - 1 == 5

    t6_dev = doc.tables[6]
    print(f"[DOCX] Developer Backlog items: {len(t6_dev.rows) - 1} (Expected: 11)")
    assert len(t6_dev.rows) - 1 == 11

    t7_qa = doc.tables[7]
    print(f"[DOCX] QA Engineer Backlog items: {len(t7_qa.rows) - 1} (Expected: 7)")
    assert len(t7_qa.rows) - 1 == 7

    t8_ai = doc.tables[8]
    print(f"[DOCX] AI Evaluator Backlog items: {len(t8_ai.rows) - 1} (Expected: 8)")
    assert len(t8_ai.rows) - 1 == 8

    t9_constraints = doc.tables[9]
    print(f"[DOCX] Constraints rows: {len(t9_constraints.rows) - 1} (Expected: 7)")
    assert len(t9_constraints.rows) - 1 == 7

    t10_stakeholders = doc.tables[10]
    print(f"[DOCX] Stakeholders rows: {len(t10_stakeholders.rows) - 1} (Expected: 9)")
    assert len(t10_stakeholders.rows) - 1 == 9

    print("\n>>> ALL DOCX CHECKS PASSED SUCCESSFULLY! <<<\n")

    # 2. Verify Excel Spreadsheet
    excel_path = 'temp/C1SE.20_ProductBacklog-UserStory_Storelens_v1.1.xlsx'
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['Product Backlog']

    storelens_count = 0
    items_count = 0
    sprint_hours = {}

    for r in range(8, ws.max_row + 1):
        id_val = ws.cell(row=r, column=2).value
        if id_val and id_val.startswith('PB-'):
            items_count += 1
            sprint = ws.cell(row=r, column=13).value
            hours = int(ws.cell(row=r, column=14).value or 0)
            sprint_hours[sprint] = sprint_hours.get(sprint, 0) + hours
            for c in range(1, 16):
                val = str(ws.cell(row=r, column=c).value or '')
                if any(k in val.lower() for k in ['storelens', 'camera', 'rtsp', 'yolov8', 'roi', 'dwell time']):
                    storelens_count += 1

    print(f"[EXCEL] Total PB items: {items_count} (Expected: 30)")
    assert items_count == 30, f"Expected 30 items, got {items_count}"

    print(f"[EXCEL] Storelens residual text leaks: {storelens_count} (Expected: 0)")
    assert storelens_count == 0, "Error: Storelens keywords found in Excel!"

    expected_sprint_hours = {
        'Sprint 1': 336,
        'Sprint 2': 336,
        'Sprint 3': 336,
        'Sprint 4': 144,
        'Sprint 5': 336,
        'Sprint 6': 168,
        'Sprint 7': 264
    }

    print("[EXCEL] Sprint Hours Validation:")
    for sprint, expected_h in expected_sprint_hours.items():
        actual_h = sprint_hours.get(sprint, 0)
        print(f"  {sprint}: {actual_h} hours (Expected: {expected_h})")
        assert actual_h == expected_h, f"Hour mismatch for {sprint}: actual {actual_h} vs expected {expected_h}"

    total_hours = sum(sprint_hours.values())
    print(f"[EXCEL] Total Development Hours: {total_hours} (Expected: 1920)")
    assert total_hours == 1920

    print("\n>>> ALL EXCEL CHECKS PASSED SUCCESSFULLY! <<<\n")
    print("=" * 60)
    print("ALL VERIFICATIONS COMPLETED SUCCESSFULLY WITH 100% PASS RATE!")
    print("=" * 60)

if __name__ == '__main__':
    verify_all()
