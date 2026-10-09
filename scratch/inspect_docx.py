import docx

doc = docx.Document('temp/C1SE.30_ProductBacklog_CovAI_ver1.1.docx')

print("Paragraph count:", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        runs_info = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold, r.italic) for r in p.runs]
        print(f"P{i} [{p.style.name}]: {p.text[:60]}... | Runs: {len(p.runs)}")

print("\n--- Tables ---")
for t_idx, tbl in enumerate(doc.tables):
    print(f"Table {t_idx}: {len(tbl.rows)} rows x {len(tbl.columns)} cols")
    # Check header font and cell font
    if len(tbl.rows) > 0:
        h_cell = tbl.rows[0].cells[0]
        h_runs = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold) for p in h_cell.paragraphs for r in p.runs]
        print(f"  Header cell (0,0): {h_runs}")
    if len(tbl.rows) > 1:
        d_cell = tbl.rows[1].cells[0]
        d_runs = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold) for p in d_cell.paragraphs for r in p.runs]
        print(f"  Data cell (1,0): {d_runs}")
